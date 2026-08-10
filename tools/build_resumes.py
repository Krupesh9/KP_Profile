from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resume"
OUT.mkdir(exist_ok=True)

NAVY = "102A43"
BLUE = "146C94"
CYAN = "19A7CE"
TEAL = "0F766E"
INK = "172B4D"
MUTED = "52647A"
LIGHT = "EDF5F7"
PALE = "F5F8FA"
WHITE = "FFFFFF"
GOLD = "D29B2A"


PROFILE = (
    "AI and Power Platform solutions architect with 15+ years of hands-on experience across SharePoint, "
    "Microsoft 365, enterprise workflow, RPA, cloud integration, and modern AI. I connect agents and copilots "
    "to real business systems through MCP, APIs, Azure, and Power Platform, then add the governance, CI/CD, "
    "observability, security, and reusable engineering practices teams need to ship dependable solutions in "
    "weeks instead of months."
)

CORE = [
    "Power Platform architecture",
    "Power Apps & Power Automate",
    "SharePoint & Microsoft 365",
    "RPA & workflow modernization",
    "Agentic AI & orchestration",
    "Model Context Protocol (MCP)",
    "Copilot Studio & AI Builder",
    "Azure Functions & Logic Apps",
    "Azure DevOps & CI/CD",
    "Enterprise API integration",
    "Power Platform governance",
    "DLP, security & permissions",
    "Python, C#, JavaScript & SQL",
    "AI-assisted engineering",
]

EXPERIENCE = [
    {
        "company": "Hunt Oil Company",
        "location": "Dallas, TX",
        "period": "Apr 2023 - Present",
        "title": "Senior Solution Architect | AI, Agents, MCP, Azure & Power Platform",
        "bullets": [
            "Built an enterprise MCP SDK and onboarding pattern that brought 10+ MCP servers into a consistent, supportable ecosystem, including ServiceNow, Tanium, SharePoint Lists, Microsoft 365 reporting/licensing, Rapid7, Mimecast, SailPoint, Nutanix, and Palo Alto Panorama.",
            "Designed agent tools and orchestration patterns that connect reasoning workflows to real enterprise systems while keeping authentication, permissions, logging, retries, and human approval points explicit.",
            "Created Agent SDK and Agent Orchestrator capabilities for employee onboarding scenarios, including profile/persona cloning, task creation, access coordination, progress monitoring, and reusable onboarding knowledge.",
            "Developing autonomous operations agents that start from a daily task list, monitor work, detect failures, attempt safe self-healing, and build knowledge from completed work and operational feedback.",
            "Used Power Platform MCP, Copilot Studio, AI prompts, Power Automate, and Power Apps to generate reports and working automations faster - helping move delivery cycles from months toward weeks.",
            "Established Azure DevOps projects, repository and environment permissions, automated CI/CD pipelines, solution promotion, and release controls for AI and Power Platform delivery.",
            "Built mobile-friendly Power Apps and code apps with agentic AI experiences, reusable responsive components, adaptive cards, and a unified UI that reduces hand-built work across projects.",
            "Set up Power Platform governance with environment strategy, DLP policies, security roles, connector controls, and deployment standards to balance innovation with enterprise risk management.",
            "Created reusable Claude and ChatGPT/Codex skills, starter patterns, and review checklists so junior developers and interns can contribute safely to Power Platform automation, adaptive cards, and standardized app UI.",
            "Used Claude Code and ChatGPT Codex as engineering partners to draft and review C#, JavaScript, Python, Power Platform, and UI changes while keeping final technical decisions and code quality under human review.",
            "Previously delivered enterprise MDM/vendor onboarding, SAP S/4HANA and OpenInvoice integration, DocuSign automation, IT portfolio management, and a 30+ process UiPath-to-Power-Automate migration completed in under two months with $40K+ estimated annual savings.",
        ],
    },
    {
        "company": "Wells Fargo Bank",
        "location": "Dallas, TX",
        "period": "Dec 2022 - Oct 2023",
        "title": "Lead Application Architect | M365, Power Platform & Migration",
        "bullets": [
            "Led governance, information architecture, and migration delivery for 1,200+ SharePoint 2013 sites moving to SharePoint Online.",
            "Built an automated site-intake and approval experience with Power Apps and Power Automate, backed by reporting for leadership visibility.",
            "Led offshore delivery, scheduled migrations with ShareGate/AvePoint, and modernized shared navigation and UI through SPFx solutions.",
            "Designed Power Apps against large datasets of 200K+ records using Dataverse, SQL, and SharePoint patterns appropriate to scale.",
        ],
    },
    {
        "company": "Live Nation / Ticketmaster",
        "location": "Dallas, TX",
        "period": "Oct 2022 - Mar 2023",
        "title": "Lead Application Architect | Supplier Automation",
        "bullets": [
            "Architected a supplier-intake solution with tiered approvals, DocuSign, Dataverse, Power Apps, Power Automate, and Power BI.",
            "Automated extraction of DocuSign template data into Dataverse and created operational reporting for vendor status and adoption.",
            "Mentored a three-person team on Power Platform fundamentals, solution design, reusable patterns, and release automation.",
        ],
    },
    {
        "company": "Bravo Tech / Hunt Consolidated",
        "location": "Dallas, TX",
        "period": "Oct 2021 - Dec 2022",
        "title": "Lead Application Developer | Automation & Power Platform",
        "bullets": [
            "Delivered UiPath automation for utility-statement extraction and document workflows spanning DocuSign, PDF data, SQL, SharePoint, and Power Platform.",
            "Built compliance and contractor-management applications and translated stakeholder needs into Figma prototypes and production solutions.",
            "Modernized and retired legacy K2 applications while documenting supportable replacement patterns.",
        ],
    },
    {
        "company": "American Airlines",
        "location": "Fort Worth, TX",
        "period": "Jun 2018 - Oct 2021",
        "title": "Lead Application Developer | Enterprise UX & Automation",
        "bullets": [
            "Led UI architecture and delivery across enterprise K2 and SharePoint products, creating reusable responsive themes and interaction patterns.",
            "Built a Flight Operations messaging application used to prevent overfly events; the program was credited internally with approximately $2B in annual savings.",
            "Delivered a complex aircraft-parts setup system used by 20+ groups and integrated SQL, Teradata, Alteryx, K2, and web technologies.",
            "Created digital-signature controls and responsive workflows supporting FAA-governed operational processes.",
        ],
    },
]

EARLIER = [
    ("Range Resources", "Sr. SharePoint Developer", "Dec 2015 - Jul 2017", "Mobile-first Nintex forms, workflows, signatures, geolocation, and SharePoint intranet delivery."),
    ("Fannie Mae", "Sr. SharePoint Architect / Developer / UI Developer", "Aug 2014 - Dec 2015", "SP2007-to-SP2013 migration, approval workflows, responsive UI, and iPad check-in applications."),
    ("Institute for Intergovernmental Research (IIR/DHS)", "Sr. SharePoint Developer / Administrator / Architect", "Aug 2013 - Jul 2014", "Owned NIEM.gov enhancements, hybrid SharePoint administration, search, ADFS integration, and custom solutions."),
    ("Bluetooth SIG", "Sr. Software Developer", "Feb 2013 - Jul 2013", "Redesigned bluetooth.org and delivered SharePoint architecture, deployment automation, navigation, voting, and attendance features."),
    ("Lazard", "Software Application Developer", "Apr 2012 - Feb 2013", "Built secure global document-management, workflow, metadata, and SharePoint deployment solutions."),
    ("Ernst & Young", "Software Developer", "Jul 2011 - Apr 2012", "Implemented SharePoint document automation, ribbon customization, workflows, and publishing controls."),
]

CERTS = [
    "Microsoft Certified: Power Platform Solution Architect Expert (2024)",
    "Microsoft Certified: Power Platform Developer Associate",
    "UiPath Certified RPA Developer (2022)",
    "Nintex Workflow Pro Certified | K2 Certified Architect (2021)",
    "Microsoft SharePoint Developer certifications (SharePoint 2010/2013)",
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 0) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=10, bold=False, color=INK, name="Aptos", italic=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def add_hyperlink(paragraph, text: str, url: str, size=9.5, color=BLUE, name="Aptos") -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    r_pr.extend([r_fonts, color_node, size_node])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, 8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_node, end])


def add_bullet_numbering(document: Document, color=BLUE) -> int:
    numbering = document.part.numbering_part.element
    existing_abstract = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(existing_abstract or [0]) + 1
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    r_pr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    r_pr.append(col)
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id: int, compact=False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    n_id = OxmlElement("w:numId")
    n_id.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, n_id])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(2.5 if compact else 4)
    paragraph.paragraph_format.line_spacing = 1.08 if compact else 1.12


def style_document(doc: Document, compact=False) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72 if not compact else 0.42)
    section.bottom_margin = Inches(0.68 if not compact else 0.42)
    section.left_margin = Inches(0.72 if not compact else 0.42)
    section.right_margin = Inches(0.72 if not compact else 0.42)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10 if not compact else 8.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(4 if not compact else 2)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, before, after in (
        ("Heading 1", 14.5 if not compact else 11.5, 10 if not compact else 5, 4),
        ("Heading 2", 11.5 if not compact else 9.5, 7 if not compact else 3, 2),
        ("Heading 3", 10.5 if not compact else 8.8, 4, 2),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(BLUE if style_name != "Heading 3" else NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_section_heading(doc, text: str) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text.upper())
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), CYAN)
    border.append(bottom)
    p_pr.append(border)


def add_job(doc, job: dict, num_id: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(f"{job['company']} | {job['location']}")
    set_run(run, 11, bold=True, color=NAVY, name="Aptos Display")
    p.add_run("\t")
    run = p.add_run(job["period"])
    set_run(run, 9.5, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(job["title"])
    set_run(run, 10, bold=True, color=TEAL)

    for bullet in job["bullets"]:
        p = doc.add_paragraph()
        apply_bullet(p, num_id)
        set_run(p.add_run(bullet), 9.6, color=INK)


def build_full_resume() -> Path:
    doc = Document()
    style_document(doc)
    doc.core_properties.title = "Krupesh Patel - AI and Power Platform Solutions Engineer / Architect Resume"
    doc.core_properties.subject = "ATS-ready resume for AI, Power Platform, automation, and enterprise solution architecture roles"
    doc.core_properties.keywords = ", ".join(CORE[:9])[:255]

    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("KRUPESH PATEL  |  AI & POWER PLATFORM SOLUTIONS ARCHITECT"), 8.2, bold=True, color=MUTED)
    footer = section.footer
    add_page_number(footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("KRUPESH PATEL"), 25, bold=True, color=NAVY, name="Aptos Display")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("AI & Power Platform Solutions Engineer  |  Enterprise Architect"), 13, bold=True, color=BLUE, name="Aptos Display")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("Dallas-Fort Worth, TX  |  (251) 581-4769  |  "), 9.5, color=MUTED)
    add_hyperlink(p, "kp-9287@hotmail.com", "mailto:kp-9287@hotmail.com")
    set_run(p.add_run("  |  "), 9.5, color=MUTED)
    add_hyperlink(p, "linkedin.com/in/krupesh87", "https://www.linkedin.com/in/krupesh87")
    set_run(p.add_run("  |  Green Card EAD"), 9.5, color=MUTED)

    add_section_heading(doc, "Professional Profile")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(PROFILE), 10, color=INK)

    add_section_heading(doc, "Core Expertise")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run("  |  ".join(CORE)), 9.4, bold=True, color=INK)

    add_section_heading(doc, "Selected Enterprise Impact")
    impact_num = add_bullet_numbering(doc, CYAN)
    impacts = [
        "Built on 15+ years and 125+ completed projects across SharePoint, Microsoft 365, workflow, and Power Platform to create secure modernization roadmaps that teams can operate and extend.",
        "Created a reusable MCP SDK and onboarded 10+ enterprise MCP servers across ITSM, security, infrastructure, identity, collaboration, and Microsoft 365 reporting.",
        "Accelerated Power Platform and AI delivery with governed DevOps, reusable components, MCP-enabled generation, and AI-assisted engineering - moving practical solutions from months toward weeks.",
        "Migrated 30+ UiPath automations to Power Automate RPA in under two months, with $40K+ estimated annual savings and stronger queue-based auditability.",
        "Migrated and modernized 1,200+ SharePoint sites and delivered enterprise applications used across 20+ operating groups.",
    ]
    for text in impacts:
        p = doc.add_paragraph()
        apply_bullet(p, impact_num)
        set_run(p.add_run(text), 9.6)

    add_section_heading(doc, "Professional Experience")
    job_num = add_bullet_numbering(doc, CYAN)
    for job in EXPERIENCE:
        add_job(doc, job, job_num)

    add_section_heading(doc, "Earlier Experience")
    for company, title, period, summary in EARLIER:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
        set_run(p.add_run(company), 10.2, bold=True, color=NAVY)
        p.add_run("\t")
        set_run(p.add_run(period), 9, bold=True, color=MUTED)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(f"{title} - "), 9.2, bold=True, color=TEAL)
        set_run(p.add_run(summary), 9.2, color=INK)

    add_section_heading(doc, "Certifications & Education")
    cert_num = add_bullet_numbering(doc, CYAN)
    for cert in CERTS:
        p = doc.add_paragraph()
        apply_bullet(p, cert_num, compact=True)
        set_run(p.add_run(cert), 9.4)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("M.S., Computer Science"), 9.8, bold=True, color=NAVY)
    set_run(p.add_run(" - University of Bridgeport | GPA 3.76/4.0"), 9.4, color=INK)
    p = doc.add_paragraph()
    set_run(p.add_run("B.E., Information Technology"), 9.8, bold=True, color=NAVY)
    set_run(p.add_run(" - BVM Engineering College, S.P. University"), 9.4, color=INK)

    path = OUT / "Krupesh_Patel_AI_and_Power_Platform_Resume_V2.docx"
    doc.save(path)
    return path


def add_compact_heading(cell, text: str, color=CYAN) -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(text.upper()), 8.5, bold=True, color=color, name="Aptos Display")


def add_rating(cell, label: str, score: float) -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5)
    set_run(p.add_run(label), 7.7, bold=True, color=WHITE)
    p.add_run("\n")
    score = float(score)
    whole = int(score)
    half = score - whole >= 0.5
    marks = "●" * whole + ("◐" if half else "") + "○" * (5 - whole - int(half))
    label_score = str(int(score)) if score.is_integer() else str(score)
    set_run(p.add_run(f"{marks}  {label_score}/5"), 7.4, color="7DD3FC")


def add_sidebar_line(cell, text: str, bold=False) -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(text), 7.5, bold=bold, color=WHITE)


def add_compact_bullet(cell, text: str, num_id: int, size=8.0) -> None:
    p = cell.add_paragraph()
    apply_bullet(p, num_id, compact=True)
    set_run(p.add_run(text), size, color=INK)


def build_one_page_resume() -> Path:
    doc = Document()
    style_document(doc, compact=True)
    doc.core_properties.title = "Krupesh Patel - Modern One Page AI and Power Platform Architect Profile"
    doc.core_properties.subject = "One-page visual networking resume with an ATS-readable detailed-resume companion"
    doc.core_properties.keywords = ", ".join(CORE[:9])[:255]
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    section.footer_distance = Inches(0.18)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("KRUPESH PATEL  |  AI + POWER PLATFORM ARCHITECT  |  DALLAS-FORT WORTH"), 6.6, bold=True, color=MUTED)

    hero = doc.add_table(rows=1, cols=2)
    hero.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(hero, [7900, 3120], indent_dxa=0)
    remove_table_borders(hero)
    hero_left, hero_right = hero.rows[0].cells
    set_cell_fill(hero_left, NAVY)
    set_cell_fill(hero_right, NAVY)
    set_cell_margins(hero_left, 165, 220, 130, 120)
    set_cell_margins(hero_right, 100, 120, 100, 120)
    hero_left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hero_right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p = hero_left.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("KRUPESH PATEL"), 24, bold=True, color=WHITE, name="Aptos Display")
    p = hero_left.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run("AI & POWER PLATFORM SOLUTIONS ENGINEER / ARCHITECT"), 10.1, bold=True, color="7DD3FC", name="Aptos Display")
    p = hero_left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("Microsoft 365 foundation  ·  Power Platform delivery  ·  AI-enabled future"), 7.7, bold=True, color="D7E9F5")

    color_strip = hero_left.add_table(rows=1, cols=4)
    set_table_geometry(color_strip, [1880, 1880, 1880, 1880])
    remove_table_borders(color_strip)
    for cell, fill in zip(color_strip.rows[0].cells, ["4285F4", "8B5CF6", "EC4899", "F59E0B"]):
        set_cell_fill(cell, fill)
        set_cell_margins(cell, 18, 0, 18, 0)

    p = hero_right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    photo = ROOT / "profilePhoto-Krupesh.jpeg"
    if photo.exists():
        p.add_run().add_picture(str(photo), width=Inches(1.18), height=Inches(1.48))

    body = doc.add_table(rows=1, cols=2)
    body.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(body, [3460, 7560], indent_dxa=0)
    remove_table_borders(body)
    left, right = body.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_fill(left, "EAF2F7")
    set_cell_fill(right, WHITE)
    set_cell_margins(left, top=120, start=170, bottom=105, end=170)
    set_cell_margins(right, top=120, start=235, bottom=105, end=80)

    def left_heading(text: str) -> None:
        p = left.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(text.upper()), 8.7, bold=True, color=BLUE, name="Aptos Display")

    def right_heading(text: str) -> None:
        p = right.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(text.upper()), 9.4, bold=True, color=BLUE, name="Aptos Display")

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("PROFILE"), 8.7, bold=True, color=BLUE, name="Aptos Display")
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("Architect and hands-on engineer with 15+ years since 2011 across SharePoint, Microsoft 365, Power Platform, RPA, Azure, APIs, MCP, and enterprise agents."), 7.8, color=INK)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("I help teams adopt AI as a practical capability layer—without losing governance, security, usability, or supportability."), 7.8, bold=True, color=TEAL)

    left_heading("At a glance")
    metrics = left.add_table(rows=2, cols=2)
    set_table_geometry(metrics, [1560, 1560])
    remove_table_borders(metrics)
    for cell, big, small in zip(
        [metrics.cell(0, 0), metrics.cell(0, 1), metrics.cell(1, 0), metrics.cell(1, 1)],
        ["15+", "10+", "30+", "1,200+"],
        ["years since 2011", "MCP servers", "RPA migrations", "SharePoint sites migrated"],
    ):
        set_cell_fill(cell, WHITE)
        set_cell_margins(cell, 55, 45, 55, 45)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(big), 11.5, bold=True, color=NAVY, name="Aptos Display")
        p.add_run("\n")
        set_run(p.add_run(small), 6.5, bold=True, color=MUTED)

    left_heading("Capabilities")
    for label, score in [
        ("Power Platform", 5),
        ("Power Apps / Automate / RPA", 5),
        ("SharePoint / Microsoft 365", 5),
        ("MCP / APIs / Integration", 5),
        ("Agents / Copilot Studio", 4.5),
        ("Azure / DevOps / CI/CD", 4.5),
        ("C# / JS / SQL / Python", 4.5),
        ("Power BI", 2.5),
    ]:
        p = left.add_paragraph()
        p.paragraph_format.space_after = Pt(1.1)
        set_run(p.add_run(label), 6.9, bold=True, color=INK)
        p.add_run("\n")
        score = float(score)
        whole = int(score)
        half = score - whole >= 0.5
        marks = "●" * whole + ("◐" if half else "") + "○" * (5 - whole - int(half))
        label_score = str(int(score)) if score.is_integer() else str(score)
        set_run(p.add_run(f"{marks}  {label_score}/5"), 6.8, color=BLUE)

    left_heading("Contact")
    for text in ["Dallas-Fort Worth, TX", "(251) 581-4769", "kp-9287@hotmail.com", "linkedin.com/in/krupesh87", "Green Card EAD"]:
        p = left.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(text), 7.0, color=INK)

    left_heading("Credentials")
    for text in ["Power Platform Solution Architect Expert", "Power Platform Developer Associate", "UiPath RPA Developer", "K2 Architect · Nintex Workflow Pro"]:
        p = left.add_paragraph()
        p.paragraph_format.space_after = Pt(1.2)
        set_run(p.add_run(text), 6.8, bold=True, color=INK)

    right_heading("The work I am known for")
    p = right.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("I connect Microsoft platforms, automation, and AI to the enterprise systems where work actually happens."), 10.3, bold=True, color=TEAL, italic=True, name="Aptos Display")

    right_heading("Experience")
    snapshots = [
        ("2023—NOW", "Hunt Oil Company", "Senior Solution Architect", "Built an enterprise MCP SDK and onboarded 10+ servers; delivering agent orchestration, Copilot Studio, Power Platform governance, mobile apps, and Azure DevOps CI/CD."),
        ("2022—23", "Wells Fargo", "Lead Application Architect", "Led governance and migration delivery for 1,200+ SharePoint sites plus Power Platform intake, reporting, and SPFx modernization."),
        ("2022—23", "Live Nation / Ticketmaster", "Lead Application Architect", "Delivered supplier onboarding with Power Apps, Power Automate, Dataverse, DocuSign, approvals, reporting, and team enablement."),
        ("2021—22", "BravoTECH / Hunt", "Lead Application Developer", "Built RPA, document-extraction, compliance, contractor, and legacy-modernization solutions."),
        ("2018—21", "American Airlines", "Lead Application Developer", "Led enterprise UX and regulated workflows; delivered Flight Ops and aircraft-parts applications across 20+ groups."),
        ("2011—17", "Range · Fannie Mae · IIR/DHS · Bluetooth SIG · Lazard · EY", "SharePoint Architect / Developer", "Established the SharePoint, workflow, migration, document-management, and responsive enterprise UI foundation that the current AI work builds on."),
    ]
    for period, company, title, detail in snapshots:
        timeline = right.add_table(rows=1, cols=2)
        set_table_geometry(timeline, [1100, 6240])
        remove_table_borders(timeline)
        period_cell, detail_cell = timeline.rows[0].cells
        set_cell_margins(period_cell, 25, 0, 20, 70)
        set_cell_margins(detail_cell, 25, 45, 20, 0)
        set_cell_fill(period_cell, "EAF2F7")
        p = period_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(period), 6.8, bold=True, color=BLUE)
        p = detail_cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(company), 8.2, bold=True, color=NAVY, name="Aptos Display")
        p.add_run("\n")
        set_run(p.add_run(title), 7.1, bold=True, color=TEAL)
        p.add_run("  ·  ")
        set_run(p.add_run(detail), 7.1, color=INK)

    right_heading("Selected impact")
    impact = right.add_table(rows=1, cols=3)
    set_table_geometry(impact, [2447, 2447, 2446])
    remove_table_borders(impact)
    for cell, big, detail in zip(
        impact.rows[0].cells,
        ["WEEKS", "$40K+", "20+"],
        ["Delivery cycles moved from months toward weeks", "Estimated annual savings from RPA modernization", "Operating groups supported by enterprise applications"],
    ):
        set_cell_fill(cell, "F4F8FA")
        set_cell_margins(cell, 70, 75, 70, 75)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(big), 10.5, bold=True, color=NAVY, name="Aptos Display")
        p.add_run("\n")
        set_run(p.add_run(detail), 6.6, color=MUTED)

    right_heading("Career evolution")
    evolution = right.add_table(rows=1, cols=3)
    set_table_geometry(evolution, [2447, 2447, 2446])
    remove_table_borders(evolution)
    for cell, period, title, detail, fill in zip(
        evolution.rows[0].cells,
        ["2011+", "2018+", "2023+"],
        ["MICROSOFT FOUNDATION", "AUTOMATION AT SCALE", "APPLIED ENTERPRISE AI"],
        ["SharePoint, M365, workflow, enterprise UI", "Power Platform, RPA, cloud integration", "MCP, Copilot Studio, agents, orchestration"],
        ["E8F1FF", "F1EBFF", "FFF0F7"],
    ):
        set_cell_fill(cell, fill)
        set_cell_margins(cell, 65, 70, 65, 70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(period), 7.1, bold=True, color=BLUE)
        p.add_run("\n")
        set_run(p.add_run(title), 6.5, bold=True, color=NAVY)
        p.add_run("\n")
        set_run(p.add_run(detail), 6.4, color=MUTED)

    p = right.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("EDUCATION  "), 7.3, bold=True, color=BLUE)
    set_run(p.add_run("M.S. Computer Science, University of Bridgeport (3.76/4.0)  ·  B.E. Information Technology"), 7.2, color=INK)

    path = OUT / "Krupesh_Patel_AI_and_Power_Platform_One_Page_V2.docx"
    doc.save(path)
    return path


def build_one_page_resume_editorial() -> Path:
    """Build the redesigned V2 one-page editorial profile; legacy V2 builder remains for rollback."""
    doc = Document()
    style_document(doc, compact=True)
    doc.core_properties.title = "Krupesh Patel - One Page AI and Power Platform Solutions Architect Profile"
    doc.core_properties.subject = "Editorial one-page resume for AI, Power Platform, automation, and enterprise architecture roles"
    doc.core_properties.keywords = ", ".join(CORE[:9])[:255]

    # compact_reference_guide with a named "editorial profile sheet" override.
    section = doc.sections[0]
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.15)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(7.8)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = 1.05

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("KRUPESH PATEL  ·  AI + POWER PLATFORM SOLUTIONS ARCHITECT  ·  DALLAS–FORT WORTH"), 6.1, bold=True, color=MUTED, name="Arial")

    def gradient_text(paragraph, text: str, size: float) -> None:
        stops = [(66, 133, 244), (139, 92, 246), (236, 72, 153), (245, 158, 11)]
        visible = max(len(text) - 1, 1)
        for index, character in enumerate(text):
            position = index / visible * (len(stops) - 1)
            left = min(int(position), len(stops) - 2)
            amount = position - left
            start, end = stops[left], stops[left + 1]
            color = "".join(f"{round(start[channel] + (end[channel] - start[channel]) * amount):02X}" for channel in range(3))
            set_run(paragraph.add_run(character), size, bold=True, color=color, name="Arial")

    def add_label(cell, text: str, before=4) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1)
        gradient_text(p, text.upper(), 8.5)

    def add_role(cell, company: str, title: str, period: str, detail: str) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(2.2)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.1), WD_TAB_ALIGNMENT.RIGHT)
        set_run(p.add_run(company), 7.35, bold=True, color=NAVY, name="Arial")
        p.add_run("\t")
        set_run(p.add_run(period), 6.1, bold=True, color=BLUE, name="Arial")
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(.7)
        set_run(p.add_run(title), 6.5, bold=True, color=TEAL, name="Arial")
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.2)
        p.paragraph_format.line_spacing = 1.0
        set_run(p.add_run(detail), 6.25, color=MUTED, name="Arial")

    def add_project(cell, title: str, subtitle: str, bullets: list[str], num_id: int) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(2.4)
        p.paragraph_format.space_after = Pt(.4)
        gradient_text(p, title, 8.1)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(.8)
        set_run(p.add_run(subtitle), 6.6, bold=True, color=TEAL, name="Arial")
        for bullet in bullets:
            p = cell.add_paragraph()
            apply_bullet(p, num_id, compact=True)
            p.paragraph_format.space_after = Pt(.6)
            p.paragraph_format.line_spacing = 1.0
            set_run(p.add_run(bullet), 6.8, color=INK, name="Arial")

    hero = doc.add_table(rows=1, cols=2)
    hero.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(hero, [2780, 8450], indent_dxa=0)
    remove_table_borders(hero)
    photo_cell, hero_copy = hero.rows[0].cells
    for cell in (photo_cell, hero_copy):
        set_cell_fill(cell, "0B1427")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(photo_cell, 135, 175, 115, 95)
    set_cell_margins(hero_copy, 125, 95, 110, 210)

    photo_p = photo_cell.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.paragraph_format.space_after = Pt(0)
    photo = ROOT / "profilePhoto-Krupesh.jpeg"
    if photo.exists():
        picture = photo_p.add_run().add_picture(str(photo), width=Inches(1.42), height=Inches(1.78))
        try:
            picture._inline.graphic.graphicData.pic.spPr.prstGeom.set("prst", "ellipse")
            outline = OxmlElement("a:ln")
            outline.set("w", "38100")
            solid = OxmlElement("a:solidFill")
            color = OxmlElement("a:srgbClr")
            color.set("val", "8B5CF6")
            solid.append(color)
            outline.append(solid)
            picture._inline.graphic.graphicData.pic.spPr.append(outline)
        except AttributeError:
            pass

    p = hero_copy.paragraphs[0]
    p.paragraph_format.space_after = Pt(.5)
    set_run(p.add_run("ENTERPRISE ARCHITECTURE  ·  APPLIED AI  ·  MICROSOFT CLOUD"), 6.3, bold=True, color="9FC5FF", name="Arial")
    p = hero_copy.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    gradient_text(p, "KRUPESH PATEL", 24)
    p = hero_copy.add_paragraph()
    p.paragraph_format.space_after = Pt(3.3)
    gradient_text(p, "AI & POWER PLATFORM SOLUTIONS ENGINEER / ARCHITECT", 9.6)
    p = hero_copy.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("I connect AI agents, automation, and Microsoft platforms to the enterprise systems where work actually happens—then make the solution secure, governable, and supportable."), 7.05, color="D9E4F4", name="Arial")

    contact = hero_copy.add_table(rows=1, cols=3)
    set_table_geometry(contact, [2580, 2580, 2580], indent_dxa=0)
    remove_table_borders(contact)
    contact_items = [
        ("PHONE", "(251) 581-4769", "18263A"),
        ("LINKEDIN", "linkedin.com/in/krupesh87", "18263A"),
        ("WORK AUTHORIZATION", "GREEN CARD EAD", "146B4E"),
    ]
    for cell, (label, value, fill_color) in zip(contact.rows[0].cells, contact_items):
        set_cell_fill(cell, fill_color)
        set_cell_margins(cell, 45, 55, 45, 55)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), 5.2, bold=True, color="9FB0CA", name="Arial")
        p.add_run("\n")
        set_run(p.add_run(value), 6.15, bold=True, color=WHITE, name="Arial")

    metrics = doc.add_table(rows=1, cols=4)
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(metrics, [2808, 2808, 2808, 2808], indent_dxa=0)
    remove_table_borders(metrics)
    for cell, big, detail, fill_color in zip(
        metrics.rows[0].cells,
        ["15+", "10+", "30+", "1,200+"],
        ["YEARS SINCE 2011", "ENTERPRISE MCP SERVERS", "RPA MIGRATIONS", "SHAREPOINT SITES MIGRATED"],
        ["15355A", "243A73", "452E79", "6D2D6D"],
    ):
        set_cell_fill(cell, fill_color)
        set_cell_margins(cell, 60, 50, 55, 50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(big), 11.4, bold=True, color=WHITE, name="Arial")
        p.add_run("\n")
        set_run(p.add_run(detail), 5.3, bold=True, color="D7E5F5", name="Arial")

    arc = doc.add_table(rows=1, cols=3)
    arc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(arc, [3744, 3744, 3744], indent_dxa=0)
    remove_table_borders(arc)
    for cell, year, title, detail, fill_color in zip(
        arc.rows[0].cells,
        ["2011+", "2018+", "2023+"],
        ["MICROSOFT FOUNDATION", "AUTOMATION AT SCALE", "APPLIED ENTERPRISE AI"],
        ["SharePoint · M365 · workflow · enterprise UI", "Power Platform · RPA · cloud integration", "MCP · Copilot Studio · agents · orchestration"],
        ["EAF2FF", "F1ECFF", "FFF0F7"],
    ):
        set_cell_fill(cell, fill_color)
        set_cell_margins(cell, 42, 58, 42, 58)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(f"{year}  {title}"), 5.9, bold=True, color=NAVY, name="Arial")
        p.add_run("\n")
        set_run(p.add_run(detail), 5.8, color=MUTED, name="Arial")

    body = doc.add_table(rows=1, cols=2)
    body.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(body, [6270, 4960], indent_dxa=0)
    remove_table_borders(body)
    projects_cell, career_cell = body.rows[0].cells
    projects_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    career_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_fill(projects_cell, WHITE)
    set_cell_fill(career_cell, "F5F8FB")
    set_cell_margins(projects_cell, 115, 95, 100, 180)
    set_cell_margins(career_cell, 115, 175, 100, 95)

    p = projects_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(.5)
    gradient_text(p, "SIGNATURE WORK", 10)
    project_num = add_bullet_numbering(doc, "8B5CF6")
    add_project(projects_cell, "Enterprise MCP & Agent Ecosystem", "Hunt Oil Company · 2023–Present", [
        "Built a reusable MCP SDK and onboarded 10+ servers spanning ServiceNow, Tanium, SharePoint, Microsoft reporting/licensing, Rapid7, Mimecast, SailPoint, Nutanix, and Panorama.",
        "Designed agent orchestration for employee onboarding, daily task execution, monitoring, safe self-healing, reusable knowledge, and human approval points.",
    ], project_num)
    add_project(projects_cell, "Governed Power Platform Delivery", "Power Platform · Copilot Studio · Azure DevOps", [
        "Combined Power Platform MCP, AI prompts, Power Apps, Power Automate, code apps, and responsive UI patterns to move practical delivery from months toward weeks.",
        "Established DLP, security roles, connector controls, CI/CD, permissions, reusable components, adaptive cards, and AI-assisted development standards.",
    ], project_num)
    add_project(projects_cell, "Automation & M365 Modernization", "Enterprise portfolio impact", [
        "Migrated 30+ UiPath processes to Power Automate RPA in under two months, with stronger queue operations and $40K+ estimated annual savings.",
        "Migrated and modernized 1,200+ SharePoint sites and delivered operational applications across 20+ business groups.",
    ], project_num)
    add_project(projects_cell, "Master Data & Vendor Onboarding", "Power Platform · SAP S/4HANA · DocuSign", [
        "Delivered an end-to-end vendor experience spanning validation, approvals, documents, signatures, and automated ERP vendor creation.",
        "Connected IRS and bank validation services, OpenInvoice, SAP S/4HANA, Dataverse, DocuSign, and mobile-friendly Power Apps.",
    ], project_num)
    add_project(projects_cell, "M365 Governance & Migration", "Wells Fargo · SharePoint Online · SPFx", [
        "Delivered information architecture, governance, migration execution, and leadership reporting for a 1,200+ site modernization program.",
        "Built automated intake and approval with Power Apps and Power Automate while coordinating migrations, offshore delivery, and SPFx modernization.",
    ], project_num)
    add_project(projects_cell, "Operational Applications at Airline Scale", "American Airlines · Enterprise UX · Regulated workflow", [
        "Delivered FAA-governed messaging and signature workflows plus an aircraft-parts setup system used across more than 20 operating groups.",
        "Built a Flight Operations application credited internally with approximately $2B in annual impact and established reusable responsive UI patterns.",
    ], project_num)

    p = career_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(.5)
    gradient_text(p, "EXPERIENCE", 10)
    add_role(career_cell, "Hunt Oil Company", "Senior Solution Architect", "2023–NOW", "AI/MCP architecture, agents, Copilot Studio, Power Platform, governance, CI/CD, APIs, and mobile app delivery.")
    add_role(career_cell, "Wells Fargo", "Lead Application Architect", "2022–23", "Governance and migration delivery for 1,200+ SharePoint sites, intake automation, reporting, and SPFx.")
    add_role(career_cell, "Live Nation / Ticketmaster", "Lead Application Architect", "2022–23", "Supplier onboarding, approvals, Dataverse, DocuSign, Power Apps, Power Automate, Power BI, and mentoring.")
    add_role(career_cell, "BravoTECH / Hunt", "Lead Application Developer", "2021–22", "RPA, document extraction, compliance applications, K2 modernization, and Power Platform delivery.")
    add_role(career_cell, "American Airlines", "Lead Application Developer", "2018–21", "Enterprise UX, FAA-governed workflows, and operational applications serving 20+ groups.")
    add_role(career_cell, "Range Resources", "Senior SharePoint Developer", "2015–17", "Mobile Nintex forms, signatures, geolocation, workflows, and intranet architecture.")
    add_role(career_cell, "Fannie Mae", "Senior SharePoint Architect / Developer", "2014–15", "SP2007-to-SP2013 migration, responsive UI, approvals, and iPad check-in applications.")
    add_role(career_cell, "IIR / DHS", "Senior SharePoint Architect / Administrator", "2013–14", "NIEM.gov ownership, hybrid administration, search, ADFS, workflow, and custom solutions.")
    add_role(career_cell, "Bluetooth SIG", "Senior Software Developer", "2013", "bluetooth.org redesign, SharePoint architecture, navigation, voting, and attendance.")
    add_role(career_cell, "Lazard", "Software Application Developer", "2012–13", "Global document management, metadata, secure workflow, and deployment automation.")
    add_role(career_cell, "Ernst & Young", "Software Developer", "2011–12", "Document automation, ribbon customization, workflow, and publishing controls.")

    add_label(career_cell, "Core capability", before=4)
    for label, score in [
        ("Power Platform / M365", 5), ("MCP / APIs / Integration", 5),
        ("Agents / Copilot Studio", 4.5), ("Azure DevOps / CI/CD", 4.5),
        ("C# / JS / SQL / Python", 4.5), ("Power BI", 2.5),
    ]:
        p = career_cell.add_paragraph()
        p.paragraph_format.space_after = Pt(.25)
        score_value = float(score)
        whole = int(score_value)
        half = score_value - whole >= .5
        marks = "●" * whole + ("◐" if half else "") + "○" * (5 - whole - int(half))
        set_run(p.add_run(label), 6.5, bold=True, color=INK, name="Arial")
        set_run(p.add_run(f"  {marks}"), 6.3, color=BLUE, name="Arial")

    add_label(career_cell, "Credentials & education", before=3)
    p = career_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(
        "Power Platform Solution Architect Expert · Power Platform Developer Associate · UiPath RPA Developer · K2 Architect · Nintex Workflow Pro\n"
        "M.S. Computer Science, University of Bridgeport (3.76/4.0) · B.E. Information Technology"
    ), 6.3, color=INK, name="Arial")

    ecosystem = doc.add_table(rows=1, cols=4)
    ecosystem.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(ecosystem, [2808, 2808, 2808, 2808], indent_dxa=0)
    remove_table_borders(ecosystem)
    for cell, title, detail, fill_color in zip(
        ecosystem.rows[0].cells,
        ["MCP & AGENTS", "POWER PLATFORM", "M365 & SHAREPOINT", "AZURE & DELIVERY"],
        ["Agent SDK · orchestration · Copilot Studio", "Power Apps · Automate · RPA · Dataverse", "Architecture · SPFx · migration · governance", "DevOps · CI/CD · Functions · APIs · security"],
        ["15355A", "243A73", "452E79", "6D2D6D"],
    ):
        set_cell_fill(cell, fill_color)
        set_cell_margins(cell, 360, 72, 360, 72)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(title), 6.8, bold=True, color=WHITE, name="Arial")
        p.add_run("\n")
        set_run(p.add_run(detail), 6.0, color="D7E5F5", name="Arial")

    role_fit = doc.add_table(rows=1, cols=1)
    role_fit.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(role_fit, [11232], indent_dxa=0)
    remove_table_borders(role_fit)
    set_cell_fill(role_fit.cell(0, 0), "EAF2FF")
    set_cell_margins(role_fit.cell(0, 0), 120, 120, 120, 120)
    p = role_fit.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("BEST-FIT ROLES  "), 6.6, bold=True, color=BLUE, name="Arial")
    set_run(p.add_run("AI Solutions Engineer / Architect  ·  Power Platform Solution Architect  ·  Copilot & Automation Architect  ·  M365 Modernization Architect"), 6.5, bold=True, color=NAVY, name="Arial")

    path = OUT / "Krupesh_Patel_AI_and_Power_Platform_One_Page_V2.docx"
    doc.save(path)
    return path


def build_one_page_resume_ayden() -> Path:
    """Build a clean Ayden-inspired one-page resume in the Gemini Aurora palette."""
    doc = Document()
    style_document(doc, compact=True)
    doc.core_properties.title = "Krupesh Patel - One Page AI and Power Platform Solutions Architect Resume"
    doc.core_properties.subject = "Clean one-page resume for AI, Power Platform, automation, and Microsoft cloud architecture roles"
    doc.core_properties.keywords = ", ".join(CORE[:10])[:255]

    section = doc.sections[0]
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.30)
    section.right_margin = Inches(0.30)
    section.header_distance = Inches(0.12)
    section.footer_distance = Inches(0.14)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(8.0)
    normal.paragraph_format.space_after = Pt(1.6)
    normal.paragraph_format.line_spacing = 1.04

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("KRUPESH PATEL  ·  DALLAS–FORT WORTH  ·  GREEN CARD EAD"), 6.2, bold=True, color=MUTED, name="Arial")

    def gradient_text(paragraph, text: str, size: float) -> None:
        stops = [(66, 133, 244), (139, 92, 246), (236, 72, 153), (249, 115, 22)]
        visible = max(len(text) - 1, 1)
        for index, character in enumerate(text):
            position = index / visible * (len(stops) - 1)
            left = min(int(position), len(stops) - 2)
            amount = position - left
            start, end = stops[left], stops[left + 1]
            color = "".join(f"{round(start[channel] + (end[channel] - start[channel]) * amount):02X}" for channel in range(3))
            set_run(paragraph.add_run(character), size, bold=True, color=color, name="Arial")

    def section_label(cell, text: str, color=BLUE, before=6.0, after=2.8) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        set_run(p.add_run(text.upper()), 9.1, bold=True, color=color, name="Arial")
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), "DCE5F1")
        borders.append(bottom)
        p_pr.append(borders)

    def skill_rating(cell, label: str, score: float) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1.7)
        p.paragraph_format.line_spacing = 1.0
        set_run(p.add_run(label), 7.2, bold=True, color=INK, name="Arial")
        p.add_run("\n")
        value = float(score)
        whole = int(value)
        half = value - whole >= .5
        marks = "●" * whole + ("◐" if half else "") + "○" * (5 - whole - int(half))
        set_run(p.add_run(marks), 7.4, bold=True, color=BLUE, name="Arial")
        label_score = str(int(value)) if value.is_integer() else str(value)
        set_run(p.add_run(f"   {label_score}/5"), 6.4, bold=True, color=MUTED, name="Arial")

    def role_block(cell, company: str, period: str, title: str, bullets: list[str]) -> None:
        p = cell.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(3.0)
        p.paragraph_format.space_after = Pt(.5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(4.85), WD_TAB_ALIGNMENT.RIGHT)
        set_run(p.add_run(company), 8.2, bold=True, color=NAVY, name="Arial")
        p.add_run("\t")
        set_run(p.add_run(period), 6.6, bold=True, color=BLUE, name="Arial")
        p = cell.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(.7)
        set_run(p.add_run(title), 7.2, bold=True, color=TEAL, name="Arial")
        for item in bullets:
            p = cell.add_paragraph()
            apply_bullet(p, bullet_num, compact=True)
            p.paragraph_format.space_after = Pt(.75)
            p.paragraph_format.line_spacing = 1.02
            set_run(p.add_run(item), 7.15, color=INK, name="Arial")

    accent = doc.add_table(rows=1, cols=4)
    accent.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(accent, [2808, 2808, 2808, 2808], indent_dxa=0)
    remove_table_borders(accent)
    for cell, fill in zip(accent.rows[0].cells, ["4285F4", "8B5CF6", "EC4899", "F97316"]):
        set_cell_fill(cell, fill)
        set_cell_margins(cell, 14, 0, 14, 0)

    page = doc.add_table(rows=1, cols=2)
    page.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(page, [3420, 7812], indent_dxa=0)
    remove_table_borders(page)
    rail, main = page.rows[0].cells
    rail.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_fill(rail, "EEF3FA")
    set_cell_fill(main, WHITE)
    set_cell_margins(rail, 185, 180, 150, 180)
    set_cell_margins(main, 180, 245, 150, 120)

    photo_p = rail.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_p.paragraph_format.space_after = Pt(4)
    photo = ROOT / "profilePhoto-Krupesh.jpeg"
    if photo.exists():
        photo_p.add_run().add_picture(str(photo), width=Inches(1.54), height=Inches(1.93))

    p = rail.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    gradient_text(p, "KRUPESH", 19)
    p = rail.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2.2)
    gradient_text(p, "PATEL", 19)
    p = rail.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run("AI + POWER PLATFORM\nSOLUTIONS ARCHITECT"), 8.0, bold=True, color=NAVY, name="Arial")

    section_label(rail, "About me", before=2.5)
    p = rail.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(
        "Hands-on architect with 15+ years across SharePoint, Microsoft 365, Power Platform, RPA, Azure, APIs, MCP, and enterprise agents. I turn emerging AI into secure, supportable solutions that people can use."
    ), 7.35, color=INK, name="Arial")

    section_label(rail, "Contact", before=4)
    for label, value in [
        ("LOCATION", "Dallas–Fort Worth, TX"),
        ("PHONE", "(251) 581-4769"),
        ("EMAIL", "kp-9287@hotmail.com"),
        ("LINKEDIN", "linkedin.com/in/krupesh87"),
        ("WORK AUTHORIZATION", "Green Card EAD"),
    ]:
        p = rail.add_paragraph()
        p.paragraph_format.space_after = Pt(1.65)
        set_run(p.add_run(label + "  "), 5.9, bold=True, color=BLUE, name="Arial")
        set_run(p.add_run(value), 6.9, bold=label == "WORK AUTHORIZATION", color=INK, name="Arial")

    section_label(rail, "Core skills", before=4)
    for label, score in [
        ("Power Platform", 5),
        ("Power Apps / Automate / RPA", 5),
        ("SharePoint / Microsoft 365", 5),
        ("MCP / APIs / Integration", 5),
        ("Agents / Copilot Studio", 4.5),
        ("Azure / DevOps / CI/CD", 4.5),
        ("C# / JavaScript / Python / SQL", 4.5),
        ("Power BI", 2.5),
    ]:
        skill_rating(rail, label, score)

    section_label(rail, "Education", before=4)
    p = rail.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("M.S. Computer Science"), 7.05, bold=True, color=INK, name="Arial")
    p.add_run("\n")
    set_run(p.add_run("University of Bridgeport · 3.76/4.0"), 6.55, color=MUTED, name="Arial")
    p = rail.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("B.E. Information Technology"), 7.05, bold=True, color=INK, name="Arial")

    p = main.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("AI SOLUTIONS ENGINEER / ARCHITECT  ·  POWER PLATFORM SOLUTION ARCHITECT"), 7.35, bold=True, color=BLUE, name="Arial")
    p = main.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    gradient_text(p, "ENTERPRISE AI, GROUNDED IN DELIVERY", 17)
    p = main.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(
        "I connect agents, copilots, automation, and Microsoft platforms to real enterprise systems—then add the governance, security, CI/CD, and human controls required for dependable delivery."
    ), 8.55, bold=True, color=NAVY, name="Arial")

    metrics = main.add_table(rows=1, cols=5)
    set_table_geometry(metrics, [1472, 1472, 1472, 1472, 1472])
    remove_table_borders(metrics)
    for cell, big, detail, fill in zip(
        metrics.rows[0].cells,
        ["15+", "125+", "10+", "30+", "1,200+"],
        ["YEARS", "PROJECTS COMPLETED", "MCP SERVERS", "RPA MIGRATIONS", "SHAREPOINT SITES MIGRATED"],
        ["E8F1FF", "EDF0FF", "F0EAFF", "FFF0F7", "FFF3E8"],
    ):
        set_cell_fill(cell, fill)
        set_cell_margins(cell, 70, 35, 65, 35)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(big), 10.4, bold=True, color=NAVY, name="Arial")
        p.add_run("\n")
        set_run(p.add_run(detail), 4.8, bold=True, color=MUTED, name="Arial")

    bullet_num = add_bullet_numbering(doc, "4285F4")
    section_label(main, "Professional experience", before=5)
    role_block(main, "Hunt Oil Company", "2023–Present", "Senior Solution Architect · AI, Agents, MCP, Azure & Power Platform", [
        "Built a reusable MCP SDK and onboarded 10+ enterprise servers across ITSM, security, identity, infrastructure, Microsoft 365, and business data.",
        "Deliver agent orchestration, Copilot Studio, mobile Power Apps, Power Automate, governance, DLP, and Azure DevOps CI/CD—helping move delivery from months toward weeks.",
    ])
    role_block(main, "Wells Fargo", "2022–2023", "Lead Application Architect · M365, Power Platform & Migration", [
        "Led governance and migration delivery for 1,200+ SharePoint sites; built automated intake, approvals, reporting, and SPFx modernization patterns.",
    ])
    role_block(main, "Live Nation / Ticketmaster", "2022–2023", "Lead Application Architect · Supplier Automation", [
        "Architected supplier onboarding with Power Apps, Power Automate, Dataverse, DocuSign, tiered approvals, reporting, and team enablement.",
    ])
    role_block(main, "BravoTECH / Hunt Consolidated", "2021–2022", "Lead Application Developer · Automation & Power Platform", [
        "Delivered RPA, document extraction, compliance applications, Power Platform solutions, responsive prototypes, and K2 modernization.",
    ])
    role_block(main, "American Airlines", "2018–2021", "Lead Application Developer · Enterprise UX & Automation", [
        "Led responsive enterprise UX and FAA-governed workflows; delivered Flight Operations and aircraft-parts applications across 20+ operating groups.",
    ])
    p = main.add_paragraph()
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("EARLIER EXPERIENCE  "), 6.7, bold=True, color=BLUE, name="Arial")
    set_run(p.add_run("Range Resources · Fannie Mae · IIR/DHS · Bluetooth SIG · Lazard · Ernst & Young  |  SharePoint architecture, workflow, migration, document management, and enterprise UI (2011–2017)"), 6.8, color=INK, name="Arial")

    section_label(main, "Selected project impact", before=4)
    for title, detail in [
        ("Enterprise MCP & Agent Ecosystem", "Agent SDK/orchestrator, employee onboarding, daily task execution, monitoring, safe self-healing, reusable knowledge, and human approvals."),
        ("Power Platform Delivery System", "Power Platform MCP, Copilot Studio, AI prompts, code apps, adaptive cards, reusable UI, DLP, permissions, and automated releases."),
        ("Automation Modernization", "Migrated 30+ UiPath processes to Power Automate RPA in under two months with $40K+ estimated annual savings."),
        ("Master Data & Vendor Onboarding", "Connected Power Platform, SAP S/4HANA, OpenInvoice, DocuSign, IRS/bank validation, approvals, and mobile-friendly experiences."),
    ]:
        p = main.add_paragraph()
        apply_bullet(p, bullet_num, compact=True)
        p.paragraph_format.space_after = Pt(.65)
        p.paragraph_format.line_spacing = 1.0
        set_run(p.add_run(title + " — "), 7.05, bold=True, color=NAVY, name="Arial")
        set_run(p.add_run(detail), 6.95, color=INK, name="Arial")

    section_label(main, "Credentials", before=4)
    p = main.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(
        "Microsoft Power Platform Solution Architect Expert · Power Platform Developer Associate · UiPath RPA Developer · K2 Architect · Nintex Workflow Pro · Microsoft SharePoint Developer"
    ), 6.95, color=INK, name="Arial")

    role_fit = doc.add_table(rows=1, cols=1)
    role_fit.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(role_fit, [11232], indent_dxa=0)
    remove_table_borders(role_fit)
    set_cell_fill(role_fit.cell(0, 0), "132A50")
    set_cell_margins(role_fit.cell(0, 0), 62, 90, 62, 90)
    p = role_fit.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("BEST-FIT ROLES  "), 6.3, bold=True, color="9FC5FF", name="Arial")
    set_run(p.add_run("AI Solutions Engineer / Architect  ·  Power Platform Solution Architect  ·  Copilot & Automation Architect  ·  M365 Modernization Architect"), 6.25, bold=True, color=WHITE, name="Arial")

    path = OUT / "Krupesh_Patel_AI_and_Power_Platform_One_Page_V2.docx"
    doc.save(path)
    return path


def add_placeholder(paragraph, text: str, size=9.2, bold=False) -> None:
    run = paragraph.add_run(f"[[{text}]]")
    set_run(run, size, bold=bold, color=NAVY)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF1C7")
    run._element.get_or_add_rPr().append(shd)


def build_cover_letter() -> Path:
    doc = Document()
    style_document(doc, compact=True)
    doc.core_properties.title = "Krupesh Patel - Dynamic Cover Letter Template"
    doc.core_properties.subject = "Private fill-in cover letter template for AI and Power Platform solution roles"
    doc.core_properties.keywords = "AI Solutions Architect, Power Platform Architect, MCP, Copilot Studio, Azure"

    section = doc.sections[0]
    section.footer_distance = Inches(0.18)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("PRIVATE WORKING TEMPLATE  |  CUSTOMIZE FOR EACH APPLICATION"), 6.8, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("KRUPESH PATEL"), 21, bold=True, color=NAVY, name="Aptos Display")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("AI & Power Platform Solutions Engineer / Architect"), 11.5, bold=True, color=BLUE, name="Aptos Display")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("Dallas-Fort Worth, TX  |  (251) 581-4769  |  "), 8.7, color=MUTED)
    add_hyperlink(p, "kp-9287@hotmail.com", "mailto:kp-9287@hotmail.com", size=8.7)
    set_run(p.add_run("  |  "), 8.7, color=MUTED)
    add_hyperlink(p, "linkedin.com/in/krupesh87", "https://www.linkedin.com/in/krupesh87", size=8.7)

    rule = doc.add_table(rows=1, cols=1)
    set_table_geometry(rule, [11020])
    remove_table_borders(rule)
    set_cell_fill(rule.cell(0, 0), CYAN)
    set_cell_margins(rule.cell(0, 0), 14, 0, 14, 0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("To: "), 8.7, bold=True, color=MUTED)
    add_placeholder(p, "HIRING MANAGER", 8.7, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("Company / Role: "), 8.7, bold=True, color=MUTED)
    add_placeholder(p, "TARGET COMPANY", 8.7, True)
    set_run(p.add_run("  |  "), 8.7, color=MUTED)
    add_placeholder(p, "JOB TITLE", 8.7, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    set_run(p.add_run("Date: "), 8.7, bold=True, color=MUTED)
    add_placeholder(p, "DATE", 8.7)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("Dear "), 9.5, color=INK)
    add_placeholder(p, "HIRING MANAGER", 9.5)
    set_run(p.add_run(","), 9.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("I am applying for the "), 9.4, color=INK)
    add_placeholder(p, "JOB TITLE", 9.4, True)
    set_run(p.add_run(" role at "), 9.4, color=INK)
    add_placeholder(p, "TARGET COMPANY", 9.4, True)
    set_run(p.add_run(". Your need to "), 9.4, color=INK)
    add_placeholder(p, "TOP CHALLENGE FROM THE JOB POSTING", 9.4)
    set_run(p.add_run(" stood out because it matches the work I have done across Microsoft 365, Power Platform, enterprise automation, cloud integration, and emerging AI."), 9.4, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("At Hunt Oil, I built an enterprise MCP SDK and onboarding pattern for 10+ servers spanning ServiceNow, Tanium, Microsoft 365, security, identity, and infrastructure platforms. I also combined Copilot Studio, Power Apps, Power Automate, Azure DevOps, and governed agent orchestration to help teams move useful solutions from months toward weeks. That work builds naturally on 15+ years of SharePoint, workflow, RPA, application architecture, and enterprise UX delivery."), 9.4, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("For your team, the most relevant connection is "), 9.4, color=INK)
    add_placeholder(p, "ONE MATCHING REQUIREMENT", 9.4, True)
    set_run(p.add_run(". I can contribute as both an architect and a hands-on engineer: shaping the roadmap, connecting APIs and MCP tools, building Power Platform solutions, establishing CI/CD and DLP controls, and putting human review, permissions, and observability around autonomous workflows."), 9.4, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run("I would welcome a conversation about "), 9.4, color=INK)
    add_placeholder(p, "A SPECIFIC COMPANY INITIATIVE OR OUTCOME", 9.4)
    set_run(p.add_run(" and how my background could help your team modernize responsibly while delivering visible business value. Thank you for your time and consideration."), 9.4, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run("Sincerely,"), 9.4, color=INK)
    p = doc.add_paragraph()
    set_run(p.add_run("Krupesh Patel"), 10.5, bold=True, color=NAVY, name="Aptos Display")

    note = doc.add_table(rows=1, cols=1)
    set_table_geometry(note, [11020])
    remove_table_borders(note)
    set_cell_fill(note.cell(0, 0), "FFF8E7")
    set_cell_margins(note.cell(0, 0), 75, 110, 75, 110)
    p = note.cell(0, 0).paragraphs[0]
    set_run(p.add_run("CUSTOMIZE BEFORE SENDING  "), 7.6, bold=True, color=GOLD)
    set_run(p.add_run("Replace every highlighted placeholder. Use one challenge from the posting, one matching proof point, and one company-specific reason. Keep the final letter to one page."), 7.5, color=MUTED, italic=True)

    path = OUT / "Krupesh_Patel_Dynamic_Cover_Letter_Template.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_full_resume())
    print(build_one_page_resume_ayden())
    print(build_cover_letter())
